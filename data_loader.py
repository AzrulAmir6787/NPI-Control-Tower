import io
import os

import pandas as pd
from openpyxl import load_workbook


# ============================================================
# CONFIG
# ============================================================

HEADER_ROW = 5

# Excel column E = 5
PART_COLUMN_EXCEL_INDEX = 5


# ============================================================
# FILE TYPE
# ============================================================

def _get_filename(source):
    """
    Robustly obtain filename from:
    - Streamlit UploadedFile
    - file-like object
    - adapter object
    - string path
    """

    if source is None:
        return ""

    # Streamlit UploadedFile / adapter
    for attr in [
        "name",
        "filename",
        "file_name",
        "original_name",
    ]:
        try:
            value = getattr(source, attr, None)

            if value:
                return os.path.basename(
                    str(value)
                )
        except Exception:
            pass

    # Plain path string
    if isinstance(source, str):
        return os.path.basename(source)

    return ""


def get_file_type(source):
    """
    Return file extension in lowercase.

    Example:
        .xlsx
        .xlsm
        .csv
    """

    filename = _get_filename(source)

    if not filename:
        raise ValueError(
            "Unable to determine uploaded file name."
        )

    _, extension = os.path.splitext(
        filename
    )

    extension = extension.lower().strip()

    if extension not in [
        ".xlsx",
        ".xlsm",
        ".xltx",
        ".xltm",
        ".csv",
    ]:
        raise ValueError(
            "Unsupported file type: "
            f"{extension}. "
            "Please upload an Excel "
            "(.xlsx/.xlsm) or CSV file."
        )

    return extension


# ============================================================
# FILE BYTES
# ============================================================

def get_file_bytes(source):
    """
    Convert source into raw bytes.

    Supports:
    - Streamlit UploadedFile
    - BytesIO
    - file-like objects
    - file paths
    """

    if source is None:
        raise ValueError(
            "No file supplied."
        )

    # Streamlit UploadedFile
    if hasattr(source, "getvalue"):

        try:
            return source.getvalue()
        except Exception:
            pass

    # BytesIO / file-like
    if hasattr(source, "read"):

        try:

            current_position = None

            if hasattr(source, "tell"):
                current_position = (
                    source.tell()
                )

            if hasattr(source, "seek"):
                source.seek(0)

            data = source.read()

            if (
                current_position is not None
                and hasattr(source, "seek")
            ):
                source.seek(0)

            if isinstance(data, bytes):
                return data

        except Exception:
            pass

    # Raw bytes
    if isinstance(source, bytes):
        return source

    # File path
    if isinstance(source, str):

        with open(
            source,
            "rb",
        ) as f:

            return f.read()

    raise ValueError(
        "Unable to read uploaded file."
    )


# ============================================================
# STREAM / RESET
# ============================================================

def _bytes_io(source):
    return io.BytesIO(
        get_file_bytes(source)
    )


# ============================================================
# SHEET NAMES
# ============================================================

def get_sheet_names(source):

    extension = get_file_type(
        source
    )

    if extension == ".csv":
        return ["CSV"]

    file_bytes = get_file_bytes(
        source
    )

    workbook = load_workbook(
        filename=io.BytesIO(file_bytes),
        read_only=True,
        data_only=True,
    )

    try:

        return workbook.sheetnames

    finally:

        workbook.close()


# ============================================================
# LOAD SHEET
# ============================================================

def load_sheet(
    source,
    sheet_name,
):
    """
    Load Virgin BAQ data.

    Header is row 5 in Excel,
    therefore pandas header=4.
    """

    extension = get_file_type(
        source
    )

    file_bytes = get_file_bytes(
        source
    )

    if extension == ".csv":

        return pd.read_csv(
            io.BytesIO(file_bytes),
            header=HEADER_ROW - 1,
        )

    return pd.read_excel(
        io.BytesIO(file_bytes),
        sheet_name=sheet_name,
        header=HEADER_ROW - 1,
        engine="openpyxl",
    )


# ============================================================
# EMBEDDED PICTURES
# ============================================================

def extract_embedded_pictures(
    source,
    sheet_name,
):
    """
    Extract floating / embedded Excel pictures.

    Picture is mapped to:
        Full Column E Part Number

    Returns:
        {
            "FULL-PART-NUMBER": image_bytes
        }
    """

    extension = get_file_type(
        source
    )

    # CSV cannot contain embedded Excel pictures
    if extension == ".csv":
        return {}

    file_bytes = get_file_bytes(
        source
    )

    workbook = load_workbook(
        filename=io.BytesIO(file_bytes),
        read_only=False,
        data_only=False,
    )

    picture_map = {}

    try:

        if sheet_name not in workbook.sheetnames:
            return {}

        worksheet = workbook[
            sheet_name
        ]

        # ----------------------------------------------------
        # Map Excel row number -> Column E Part Number
        # ----------------------------------------------------

        row_to_part = {}

        for row_number in range(
            HEADER_ROW + 1,
            worksheet.max_row + 1,
        ):

            cell = worksheet.cell(
                row=row_number,
                column=PART_COLUMN_EXCEL_INDEX,
            )

            value = cell.value

            if value is None:
                continue

            part = str(
                value
            ).strip()

            if not part:
                continue

            row_to_part[
                row_number
            ] = part

        # ----------------------------------------------------
        # Extract images
        # ----------------------------------------------------

        images = getattr(
            worksheet,
            "_images",
            [],
        )

        for image in images:

            try:

                anchor = image.anchor

                # openpyxl OneCellAnchor /
                # TwoCellAnchor
                from_row = (
                    anchor._from.row
                )

                excel_row = (
                    from_row + 1
                )

                part = row_to_part.get(
                    excel_row
                )

                if not part:
                    continue

                image_data = image._data()

                if not image_data:
                    continue

                # Preserve first valid picture
                # for a part.
                if part not in picture_map:

                    picture_map[
                        part
                    ] = image_data

            except Exception:
                # One bad picture should not
                # break the complete BAQ.
                continue

        return picture_map

    finally:

        workbook.close()


# ============================================================
# LOAD SHEET + PICTURES
# ============================================================

def load_sheet_with_pictures(
    source,
    sheet_name,
):

    df = load_sheet(
        source,
        sheet_name,
    )

    picture_map = (
        extract_embedded_pictures(
            source,
            sheet_name,
        )
    )

    return (
        df,
        picture_map,
    )


# ============================================================
# PICTURE LOOKUP
# ============================================================

def get_picture_for_part(
    picture_map,
    part_number,
):

    if not picture_map:
        return None

    if part_number is None:
        return None

    part = str(
        part_number
    ).strip()

    if not part:
        return None

    return picture_map.get(
        part
    )


# ============================================================
# PICTURE SUMMARY
# ============================================================

def get_picture_summary(
    picture_map,
):

    if not picture_map:

        return pd.DataFrame(
            columns=[
                "PART NUMBER",
                "PICTURE AVAILABLE",
                "IMAGE BYTES",
            ]
        )

    rows = []

    for part, image_bytes in (
        picture_map.items()
    ):

        rows.append(
            {
                "PART NUMBER": part,
                "PICTURE AVAILABLE": (
                    image_bytes
                    is not None
                ),
                "IMAGE BYTES": (
                    len(image_bytes)
                    if image_bytes
                    else 0
                ),
            }
        )

    return pd.DataFrame(
        rows
    )


# ============================================================
# COMPATIBILITY ALIASES
# ============================================================

def get_file_extension(source):
    return get_file_type(
        source
    )